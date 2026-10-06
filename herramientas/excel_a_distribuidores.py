"""
Convierte los formularios Excel de los distribuidores en el array
DISTRIBUIDORES de assets/js/distribuidores.data.js.

Uso (desde la raíz del repo):
    python herramientas/excel_a_distribuidores.py formularios/*.xlsx

Lee la hoja "Datos" de cada archivo (ignora la fila 4 de ejemplo) y
resuelve las coordenadas de cada distribuidor en este orden:
  1. Latitud / Longitud cargadas en el Excel
  2. Link de Google Maps (corto o largo): se sigue la redirección y se
     extraen las coordenadas de la URL final
  3. Geocodificación de la dirección con OpenStreetMap / Nominatim
     (gratuito, máx. 1 consulta por segundo)

Requiere: pip install openpyxl requests
"""
import sys, json, glob, re, time
from pathlib import Path
from urllib.parse import unquote
from openpyxl import load_workbook

try:
    import requests
except ImportError:
    sys.exit("Falta la librería requests:  pip install requests")

UA = "JasicArgentina-Distribuidores/1.0 (info@jasicargentina.com.ar)"
TIPO = {"distribuidor oficial": "oficial", "punto de venta": "venta",
        "servicio técnico": "servicio", "servicio tecnico": "servicio"}
def tipos_de(v):
    """'Punto de Venta, Servicio Técnico' -> ['venta','servicio']; uno solo -> 'venta'."""
    partes = re.split(r"\s*(?:,|/|\+|;|\by\b)\s*", texto(v).lower())
    t = []
    for p in partes:
        k = TIPO.get(p.strip())
        if k and k not in t:
            t.append(k)
    if not t:
        return "venta"
    return t[0] if len(t) == 1 else t

COLS = ["nombre","razon","cuit","tipo","direccion","ciudad","provincia","cp","telefono","whatsapp","email","horario",
        "web","instagram","facebook","youtube","tiktok","linkedin","maps","lat","lng","marcas","contacto","obs"]

# Rango válido para Argentina (evita errores de tipeo o lat/lng invertidas)
def en_argentina(lat, lng):
    return -56 <= lat <= -21 and -74 <= lng <= -53


# ---------------------------------------------------------------
# 1) Coordenadas desde un link de Google Maps
# ---------------------------------------------------------------
def coords_desde_link(url):
    """Devuelve (lat, lng) o None. Acepta links cortos (maps.app.goo.gl,
    goo.gl/maps) y largos (google.com/maps/...)."""
    url = (url or "").strip()
    if not url:
        return None
    try:
        # Seguir redirecciones del link corto hasta la URL larga
        r = requests.get(url, headers={"User-Agent": UA}, allow_redirects=True, timeout=15)
        final = unquote(r.url)
    except Exception as e:
        print(f"    ! no se pudo abrir el link ({e})")
        return None

    # Orden de preferencia: !3d..!4d (pin exacto) > q=lat,lng > @lat,lng (centro de la vista)
    patrones = [
        r"!3d(-?\d+\.\d+)!4d(-?\d+\.\d+)",
        r"[?&]q=(-?\d+\.\d+),(-?\d+\.\d+)",
        r"[?&]query=(-?\d+\.\d+),(-?\d+\.\d+)",
        r"/place/(-?\d+\.\d+),(-?\d+\.\d+)",
        r"@(-?\d+\.\d+),(-?\d+\.\d+)",
    ]
    # Link de "Cómo llegar": el destino viene como !1d<lng>!2d<lat> (orden invertido)
    m = re.search(r"!1d(-?\d+\.\d+)!2d(-?\d+\.\d+)", final)
    if m and en_argentina(float(m.group(2)), float(m.group(1))):
        return float(m.group(2)), float(m.group(1))
    for pat in patrones:
        m = re.search(pat, final)
        if m:
            lat, lng = float(m.group(1)), float(m.group(2))
            if en_argentina(lat, lng):
                return lat, lng
    print(f"    ! el link no contiene coordenadas reconocibles: {final[:90]}…")
    return None


# ---------------------------------------------------------------
# 2) Geocodificación de la dirección (OpenStreetMap / Nominatim)
# ---------------------------------------------------------------
_ultimo_nominatim = 0.0
def coords_desde_direccion(direccion, ciudad, provincia):
    global _ultimo_nominatim
    consulta = ", ".join(x for x in [direccion, ciudad, provincia, "Argentina"] if x)
    espera = 1.1 - (time.time() - _ultimo_nominatim)   # respetar 1 req/seg
    if espera > 0:
        time.sleep(espera)
    try:
        r = requests.get("https://nominatim.openstreetmap.org/search",
                         params={"q": consulta, "format": "json", "limit": 1, "countrycodes": "ar"},
                         headers={"User-Agent": UA}, timeout=15)
        _ultimo_nominatim = time.time()
        datos = r.json()
        if datos:
            lat, lng = float(datos[0]["lat"]), float(datos[0]["lon"])
            if en_argentina(lat, lng):
                return lat, lng
    except Exception as e:
        print(f"    ! error geocodificando ({e})")
    return None


# ---------------------------------------------------------------
# Lectura del Excel
# ---------------------------------------------------------------
def social(v, base):
    v = texto(v)
    if not v:
        return ""
    if v.startswith("http"):
        return v
    return base + v.lstrip("@").strip("/")

VACIOS = {"-", "–", "no", "no tenemos", "no tiene", "n/a", "na", "s/d", "si", "sí"}
EJEMPLO = "Ferretería Industrial Norte"   # fila de ejemplo del formulario

def texto(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    t = str(v).strip().lstrip("´'`")
    return "" if t.lower() in VACIOS else t

def leer(path):
    ws = load_workbook(path, data_only=True)["Datos"]
    out = []
    for r in ws.iter_rows(min_row=4, values_only=True):
        row = dict(zip(COLS, r))
        if texto(row["nombre"]) == EJEMPLO:
            continue
        if not texto(row["nombre"]):
            # fila suelta con solo redes sociales: se suman al distribuidor anterior
            if out:
                for k, base in (("instagram", "https://www.instagram.com/"), ("facebook", "https://www.facebook.com/")):
                    if not out[-1].get(k) and texto(row[k]):
                        out[-1][k] = social(row[k], base)
                if not out[-1].get("tiktok") and texto(row["tiktok"]):
                    out[-1]["tiktok"] = social(row["tiktok"], "https://www.tiktok.com/@")
            continue
        d = {
            "nombre":    texto(row["nombre"]),
            "tipo":      tipos_de(row["tipo"]),
            "ciudad":    texto(row["ciudad"]),
            "provincia": texto(row["provincia"]),
            "direccion": texto(row["direccion"]),
            "telefono":  texto(row["telefono"]),
            "whatsapp":  re.sub(r"\D", "", texto(row["whatsapp"])),
            "horario":   texto(row["horario"]),
            "web":       texto(row["web"]),
            "instagram": social(row["instagram"], "https://www.instagram.com/"),
            "facebook":  social(row["facebook"], "https://www.facebook.com/"),
            "tiktok":    social(row["tiktok"], "https://www.tiktok.com/@"),
            "mapsQuery": ", ".join(x for x in [texto(row["direccion"]), texto(row["ciudad"]), texto(row["provincia"])] if x),
        }
        print(f"  · {d['nombre']} ({d['ciudad']})")

        # --- Coordenadas: Excel > link de Maps > dirección ---
        coords = None
        try:
            if row["lat"] not in (None, "") and row["lng"] not in (None, ""):
                lat, lng = float(str(row["lat"]).replace(",", ".")), float(str(row["lng"]).replace(",", "."))
                if en_argentina(lat, lng):
                    coords = (lat, lng); print("    coordenadas del Excel")
                else:
                    print("    ! lat/lng fuera de Argentina, se ignoran")
        except ValueError:
            print("    ! lat/lng no numéricas, se ignoran")
        if not coords and texto(row["maps"]):
            coords = coords_desde_link(texto(row["maps"]))
            if coords: print("    coordenadas del link de Google Maps")
        if not coords and d["direccion"]:
            coords = coords_desde_direccion(d["direccion"], d["ciudad"], d["provincia"])
            if coords: print("    coordenadas geocodificadas por dirección (verificar en el mapa)")
        if coords:
            d["lat"], d["lng"] = round(coords[0], 6), round(coords[1], 6)
        else:
            print("    ! SIN COORDENADAS: aparecerá en la lista pero no en el mapa")
        out.append(d)
    return out


def main(paths):
    todos = []
    for p in paths:
        for f in sorted(glob.glob(p)):
            print("Leyendo", f)
            todos += leer(f)
    if not todos:
        sys.exit("No se encontraron filas de datos.")
    # Quitar duplicados (mismo formulario enviado dos veces)
    vistos, unicos = set(), []
    for d in todos:
        clave = (d["nombre"].lower(), d["direccion"].lower())
        if clave in vistos:
            print("  (duplicado omitido)", d["nombre"])
            continue
        vistos.add(clave); unicos.append(d)
    todos = unicos
    todos.sort(key=lambda d: (d["provincia"], d["ciudad"], d["nombre"]))
    js = ("/* Generado por herramientas/excel_a_distribuidores.py — no editar a mano */\n"
          "const DISTRIBUIDORES = " + json.dumps(todos, ensure_ascii=False, indent=2) + ";\n")
    destino = Path(__file__).resolve().parent.parent / "assets/js/distribuidores.data.js"
    destino.write_text(js, encoding="utf-8")
    sin = sum(1 for d in todos if "lat" not in d)
    print(f"\n{len(todos)} distribuidores → {destino}" + (f"  ({sin} sin coordenadas)" if sin else ""))


if __name__ == "__main__":
    main(sys.argv[1:] or ["formularios/*.xlsx"])
